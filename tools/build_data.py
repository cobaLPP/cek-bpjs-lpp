"""Olah dua file Excel BPJS (unduhan resmi) menjadi data/bpjs.json terenkripsi untuk situs Cek BPJS Pengemudi LPP.

Pemakaian:
  python3 tools/build_data.py <file_TK.xlsx> <file_KS.xlsx> data/bpjs.json [--src-tk ..] [--src-ks ..] [--periode-tk "Oktober 2026"] [--periode-ks ..]
Input boleh berupa .xlsx biasa atau hasil download Google Drive (JSON base64).
Setiap data pengemudi dienkripsi AES-GCM dengan kunci turunan NIK-nya (PBKDF2-SHA256), jadi file ini
aman dipublikasikan: isinya hanya bisa dibuka dengan NIK 16 digit yang tepat. File Excel mentah JANGAN di-commit.
"""
import sys, json, re, os, base64, hashlib, datetime, collections
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

import argparse, tempfile
ap = argparse.ArgumentParser(description="Bangun data/bpjs.json untuk Cek BPJS Pengemudi LPP")
ap.add_argument("tk"); ap.add_argument("ks"); ap.add_argument("out")
ap.add_argument("--src-tk", default=""); ap.add_argument("--src-ks", default="")
ap.add_argument("--periode-tk", default=""); ap.add_argument("--periode-ks", default="")  # isi bila periode tak terbaca dari file
A = ap.parse_args()

def as_xlsx(p):
    """Terima file .xlsx biasa ATAU hasil download_file_content Google Drive (JSON berisi base64)."""
    raw = open(p, "rb").read()
    if raw[:2] == b"PK": return p
    d = json.loads(raw.decode("utf-8"))
    if isinstance(d, list): d = json.loads(d[0]["text"]) if isinstance(d[0], dict) and "text" in d[0] else d[0]
    b = base64.b64decode(d["content"]); assert b[:2] == b"PK", "bukan file xlsx"
    f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); f.write(b); f.close(); return f.name

tk_path, ks_path, out_path = as_xlsx(A.tk), as_xlsx(A.ks), A.out
BULAN = ["Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember"]
norm = lambda v: re.sub(r"\s+", "", str(v if v is not None else "").strip().upper())
nname = lambda v: re.sub(r"\s+", " ", str(v if v is not None else "").strip().upper())
def mask(n): return " ".join(w if len(w) <= 2 else w[:2] + "*" * (len(w) - 2) for w in nname(n).split(" "))
def b64u(b): return base64.urlsafe_b64encode(b).decode().rstrip("=")
def find(h, t):
    for i, c in enumerate(h):
        if t in re.sub(r"\s+", " ", str(c or "").upper()): return i
    return -1

tk_rows, per_tk = [], ""
for ws in openpyxl.load_workbook(tk_path, read_only=True, data_only=True).worksheets:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = -1
    for i, r in enumerate(rows):
        if str(r[1] or "").strip().upper() == "PERIODE" and not per_tk:
            m = re.match(r"^(\d{1,2})/(\d{4})$", str(r[2] or ""))
            if m: per_tk = f"{BULAN[int(m[1])-1]} {m[2]}"
        if str(r[0] or "").strip() == "No.": hi = i; break
    if hi < 0: continue
    H, H2 = rows[hi], rows[hi+1]
    cR, cN, cP, cNm, cT, cC = find(H,"NOMOR REFERENSI"), find(H,"( N I K )"), find(H,"NOMOR PEGAWAI"), find(H,"NAMA TENAGA KERJA"), find(H,"TGL KEPESERTAAN"), find(H2,"CUSTOMER")
    for r in rows[hi+1:]:
        if not isinstance(r[0], (int, float)) or isinstance(r[0], bool): continue
        t = r[cT]
        tk_rows.append(dict(nik=re.sub(r"\D","",str(r[cN] or "")), kpj=str(r[cR] or "").strip(), peg=norm(r[cP]), nama=nname(r[cNm]),
            tgl=t.strftime("%Y-%m-%d") if isinstance(t, datetime.datetime) else None, cust=str(r[cC]).strip() if r[cC] else ""))

pes, fam, per_ks = [], [], ""
for ws in openpyxl.load_workbook(ks_path, read_only=True, data_only=True).worksheets:
    m = re.match(r"^(\d{4})(\d{2})", ws.title)
    if m and not per_ks: per_ks = f"{BULAN[int(m[2])-1]} {m[1]}"
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if any(str(c or "").strip().upper() == "NO JKN PEKERJA" for c in r)), -1)
    if hi < 0: continue
    H = rows[hi]
    cW, cJ, cP, cC, cH = find(H,"NO JKN PEKERJA"), find(H,"NO JKN PESERTA"), find(H,"NPP"), find(H,"CUSTOMER"), find(H,"HUBUNGAN")
    cNm = next(i for i, c in enumerate(H) if str(c or "").strip().upper() == "NAMA")
    for r in rows[hi+1:]:
        w = str(r[cW] if r[cW] is not None else "").strip()
        if not w.isdigit(): continue
        rec = dict(w=w, jkn=str(r[cJ] or "").strip(), nama=nname(r[cNm]), peg=norm(r[cP]), cust=str(r[cC]).strip() if cC >= 0 and r[cC] else "", hub=str(r[cH] or "").strip().upper())
        (pes if rec["hub"] == "PESERTA" else fam).append(rec)

recs, by_peg, by_name, bad = {}, {}, {}, []
for r in tk_rows:
    if not re.fullmatch(r"\d{16}", r["nik"]): bad.append(r); continue
    recs[r["nik"]] = {"tk": {"n": mask(r["nama"]), "kpj": r["kpj"], "tgl": r["tgl"], "cust": r["cust"]}, "ks": None}
    if r["peg"]: by_peg[r["peg"]] = r["nik"]
    by_name[r["nama"]] = "DUP" if r["nama"] in by_name else r["nik"]
fam_by = collections.defaultdict(list)
for f in fam: fam_by[f["w"]].append(f["hub"])
unmatched, matched = [], 0
for p in pes:
    nik = by_peg.get(p["peg"]) or (by_name.get(p["nama"]) if by_name.get(p["nama"]) not in (None, "DUP") else None)
    if not nik or nik not in recs: unmatched.append(p); continue
    f = fam_by.get(p["w"], []); c = collections.Counter(x.capitalize() for x in f)
    recs[nik]["ks"] = {"n": mask(p["nama"]), "jkn": p["jkn"], "cust": p["cust"], "kel": len(f), "kelD": ", ".join(f"{v} {k.lower()}" for k, v in c.items())}
    matched += 1

salt, it, out = os.urandom(16), 300000, {}
for nik, rec in recs.items():
    bits = hashlib.pbkdf2_hmac("sha256", nik.encode(), salt, it, dklen=44)
    iv = os.urandom(12)
    ct = AESGCM(bits[12:]).encrypt(iv, json.dumps(rec, separators=(",", ":"), ensure_ascii=False).encode(), None)
    out[b64u(bits[:12])] = b64u(iv) + "." + b64u(ct)

data = {"v": 1, "salt": b64u(salt), "iter": it, "recs": out, "periode": {"tk": A.periode_tk or per_tk, "ks": A.periode_ks or per_ks},
        "updatedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(), "jumlah": {"tk": len(recs), "ks": matched}, "sumber": {"tk": A.src_tk, "ks": A.src_ks}}
assert data["periode"]["tk"] and data["periode"]["ks"], "Periode tidak terbaca; isi --periode-tk / --periode-ks (mis. 'Oktober 2026')"
with open(out_path, "w", encoding="utf-8") as fh: json.dump(data, fh, separators=(",", ":"))
print(json.dumps({"periode": data["periode"], "tk": len(recs), "bad": len(bad), "ks_peserta": len(pes), "matched": matched,
    "noKS": sum(1 for r in recs.values() if not r["ks"]), "unmatched": [(p["nama"], p["peg"], p["cust"]) for p in unmatched],
    "size_kb": os.path.getsize(out_path)//1024}, ensure_ascii=False, indent=1))
